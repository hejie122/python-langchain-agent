from langchain.tools import tool
from pydantic import BaseModel, Field
from app.utils.logger import Logger
from docx import Document

#定义
class Args(BaseModel):
    path:str = Field(...,description="文件路径")
logger = Logger.get_logger(__name__)
"""
只读取word中的表格数据
"""
@tool("docx_read_tool", args_schema=Args)
def docx_read_tool(path:str)-> str:
    """
    读取docx文件
    """
    doc = Document(path)
    data = []
    # 遍历表格
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text)
            data.append(row_data)
    #返回数据
    return str(data)

if __name__ == "__main__":
    path = "C:\\Users\\33273\\Desktop\\智能体项目\\智能体项目代码\\app\\static\\upload\\word.docx"
    doc = Document(path)
    data = []
    # 遍历表格
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text)
            data.append(row_data)
    print(data)